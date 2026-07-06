'''tedtoolkit.io.xlsx - Excel (.xlsx) import/export via openpyxl.'''

import collections

import click
import openpyxl
import win32com.client

from tedtoolkit.history import _add_kwarg_to_last_command, _quoted
from tedtoolkit.prompts import ask_num, _kwarg_parse_prompt_bool
from tedtoolkit.gui.dialogs import g_sel_file, g_sel_file_to_write

ALL_SHEETS_TEXT = '**ALL SHEETS**'  # Should contain characters not allowed in an excel sheet name


def _xlsx_data_check(data):
    '''function for xlsx_export, checks to see that data are in
        the right form and raises exception if not.
        Returns single_sheet_mode variable True if data is
        list of lists instead of dict of list of lists'''
    check_ok = True
    single_sheet_mode = False
    if not isinstance(data, dict):
        if isinstance(data, list):
            if data:
                if isinstance(data[0], list):
                    single_sheet_mode = True
                else:
                    check_ok = False
            else:
                check_ok = False
        else:
            check_ok = False
    if check_ok and not single_sheet_mode:
        if not data: #if dictionary is empty
            check_ok = False
        for sheetname in data:
            if not isinstance(data[sheetname], list):
                check_ok = False
                break
            if not data[sheetname]: #if the sheet is empty:
                check_ok = False
                break
            if not isinstance(data[sheetname][0], list):
                #if first line not a list
                check_ok = False
        if check_ok:
            sheet_names = list(data.keys())
            if any([len(name) > 31 for name in sheet_names]):
                print('xlsx_export(): WARNING: Sheet names detected in excess of 31 characters; these will be truncated if possible.')
                if len(sheet_names) > len(set([name[:31] for name in sheet_names])):
                    raise Exception('xlsx_export(): Truncating long sheet names results in ambiguous names. Pass unique names with max. 31 characters')
    if not check_ok:
        raise Exception(
            'Must be list of lists table or dict of tables')
    return single_sheet_mode


def xlsx_export(data_in, **kwargs):
    '''takes in a dictionary where the keys are sheet names
        and the values are lists of lists (sheet data),
        takes optional kwarg "file_path" or, if not provided,
        prompts the user for a path via gui

        If single list of lists passed instead of dict,
        writes single table to excel.
        --kwarg "sheet_name" specifies the sheet name in this
        case'''
    if _xlsx_data_check(data_in):
        data_in = {kwargs.get('sheet_name', 'Sheet1'):data_in}
    path = kwargs.get('file_path', None)
    if path is None:
        path = g_sel_file_to_write(title='Save As XLSX (Python)', filetypes=[('XLSX', ('*.xlsx'))])
        if path[-5:] != '.xlsx':
            path = path + '.xlsx'
        _add_kwarg_to_last_command('file_path', _quoted(path), fn_name='xlsx_export')
    kwargs['wb_filter'] = _kwarg_parse_prompt_bool('wb_filter', default_val='True', prompt='Include data filter?', **kwargs)
    wb_filter = kwargs['wb_filter']
    kwargs['view_in_excel'] = _kwarg_parse_prompt_bool('view_in_excel', default_val='True', prompt='Open in Excel?', **kwargs)
    wbk = openpyxl.Workbook(write_only=not wb_filter)
    existing_sheets = [name for name in wbk.sheetnames] #should be none is not using filter
    for index, sheetname in enumerate(data_in):
        wst = wbk.create_sheet(sheetname)
        print('Exporting sheet {}...'.format(sheetname[:31]))
        with click.progressbar(data_in[sheetname],
            fill_char='>', empty_char='-') as data:
            for line in data:
                wst.append(line)
        if wb_filter:
            wst.auto_filter.ref = wst.dimensions
    if existing_sheets and not any([sheet in data_in for sheet in existing_sheets]):
        for sheet in existing_sheets:
            sheet_obj = wbk[sheet]
            wbk.remove(sheet_obj)
    print('Data prepared. Saving...')
    wbk.save(path)
    print('\nWorkbook saved under {}\n\n'.format(path))
    if kwargs.get('view_in_excel', False):
        excel = win32com.client.dynamic.Dispatch('excel.application')
        _ = excel.Workbooks.Open(path)
        excel.Visible = True
    return


def xlsx_import(**kwargs):
    '''returns dict of lists of lists with content from xlsx file
    optional keyword arguments:
    file_path (default None): full path to file or file name if
                         file is in same folder as script;
                         If no path given, a gui prompt will
                         prompt user to choose the file path
    include_formulas (default False): if True, will return
        formulas from sheet instead of values
    override_ro: overrides the (default behavior) read only
        feature; needed in some cases to fix bugs in library'''
    path = kwargs.get('file_path', None)
    ro = not kwargs.get('override_ro', False)
    if path is None:
        path = g_sel_file(title='Open XLSX (Python)',
                          filetypes=[('XLSX', ('*.xlsx'))])
    if not path:
        return None
    include_formulas = kwargs.get('include_formulas', False)
    fn = lambda line: [item.value for item in line]
    data_in = openpyxl.load_workbook(path, read_only=ro,
        data_only=(not include_formulas))
    print('Reading workbook at path: {}'.format(path))
    ret_dict = collections.OrderedDict()
    for sheet_name in data_in.sheetnames:
        print('Importing sheet {}...'.format(sheet_name))
        with click.progressbar(data_in[sheet_name],
            fill_char='*', empty_char=' ') as sheet:
            temp = [fn(line) for line in sheet]
        ret_dict[sheet_name] = temp
    return ret_dict


def ask_select_sheet(dict_in, prompt='Select a sheet to import', include_all=True):
    '''Prompt for selecting a sheet from a dict of Excel sheets, or else choosing all sheets'''
    if not isinstance(dict_in, dict):
        raise Exception('ERROR: ask_select_sheet() called without a dict of sheets as input!')
    if len(dict_in) == 1 and include_all == False:
        print(f'Only 1 sheet available: selecting it as default: {list(dict_in.keys())[0]}')
        return list(dict_in.keys())[0]
    print(prompt+'\n')
    line = '{:<5}{:35}'
    temp_dict = collections.OrderedDict()
    for index, key in enumerate(([ALL_SHEETS_TEXT] if include_all else []) + list(dict_in.keys())):
        print(line.format(index, key))
        temp_dict[index] = key
    print()
    selection = ask_num(allow_decimal=False, allow_negative=False)
    if selection not in range(len(temp_dict)+(1 if include_all else 0)):
        return ask_select_sheet(dict_in, prompt, include_all)
    return temp_dict[selection]
